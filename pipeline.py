# -*- coding: utf-8 -*-
"""Lectura y cruce de los excels: exportado de GCG, dni->proveedor y maestro
de proveedores. Devuelve, para una corrida, la lista de choferes evaluados con
la fecha de vencimiento de cada documento controlado."""
import io
import json
import re
import unicodedata
from datetime import datetime

import pandas as pd

import db
import gcg_api

SIN_JRT_LABEL = "Sin JRT / sin proveedor"


def _norm(s):
    if s is None:
        return ""
    s = str(s)
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    s = re.sub(r"[^a-z0-9]+", "", s.lower())
    return s


def _find_col(columns, *keyword_sets):
    """Busca la primer columna cuyo nombre normalizado contenga TODAS las
    palabras clave de alguno de los keyword_sets (cada set es una tupla)."""
    norm_map = {c: _norm(c) for c in columns}
    for keywords in keyword_sets:
        for col, ncol in norm_map.items():
            if all(_norm(k) in ncol for k in keywords):
                return col
    return None


def _only_digits(v):
    if v is None:
        return ""
    s = str(v)
    if s.endswith(".0"):
        s = s[:-2]
    return re.sub(r"\D", "", s)


def _parse_date(v):
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, pd.Timestamp):
        return v.date()
    if isinstance(v, datetime):
        return v.date()
    s = str(v).strip()
    if not s or s.lower() in ("nan", "nat", "-", "adeuda", "no corresponde"):
        return None
    for dayfirst in (True, False):
        try:
            return pd.to_datetime(s, dayfirst=dayfirst, errors="raise").date()
        except Exception:
            continue
    return None


# ---------------------------------------------------------------------------
# Carga de maestro de proveedores
# ---------------------------------------------------------------------------

def load_maestro_file(file_like):
    df = pd.read_excel(file_like, sheet_name=0)
    cols = df.columns.tolist()

    col_nro = _find_col(cols, ("nro", "proveedor"), ("numero", "proveedor"))
    col_razon = _find_col(cols, ("razon", "social"))
    col_jrt = _find_col(cols, ("jrt",))
    col_estado = _find_col(cols, ("estado", "actual"))
    col_tipo_serv = _find_col(cols, ("tipo", "servicio"))
    col_mail = _find_col(cols, ("mail",)) or _find_col(cols, ("email",))
    col_tel = _find_col(cols, ("telefono",))
    col_dom = _find_col(cols, ("domicilio",))
    col_prov = _find_col(cols, ("provincia",))
    col_cuit = _find_col(cols, ("cuit",))
    col_pais = _find_col(cols, ("pais",))
    col_obs = _find_col(cols, ("observaciones",))

    if col_nro is None:
        raise ValueError("No se encontró la columna 'NRO DE PROVEEDOR' en el maestro.")

    rows = []
    seen = set()
    for _, r in df.iterrows():
        nro_raw = r.get(col_nro)
        nro = _only_digits(nro_raw)
        if not nro:
            continue
        nro = int(nro)
        if nro in seen:
            continue
        seen.add(nro)
        rows.append({
            "nro_proveedor": nro,
            "razon_social": r.get(col_razon) if col_razon else None,
            "jrt": (str(r.get(col_jrt)).strip() if col_jrt and pd.notna(r.get(col_jrt)) else None),
            "estado_actual": r.get(col_estado) if col_estado else None,
            "tipo_servicio": r.get(col_tipo_serv) if col_tipo_serv else None,
            "mail": r.get(col_mail) if col_mail else None,
            "telefono": r.get(col_tel) if col_tel else None,
            "domicilio": r.get(col_dom) if col_dom else None,
            "provincia": r.get(col_prov) if col_prov else None,
            "cuit": r.get(col_cuit) if col_cuit else None,
            "pais": r.get(col_pais) if col_pais else None,
            "observaciones": r.get(col_obs) if col_obs else None,
        })
    return db.replace_proveedores(rows)


# ---------------------------------------------------------------------------
# Carga de dni -> proveedor
# ---------------------------------------------------------------------------

def load_dni_file(file_like):
    df = pd.read_excel(file_like, sheet_name=0)
    cols = df.columns.tolist()

    col_id = _find_col(cols, ("id", "vehiculo"))
    col_prop = _find_col(cols, ("prop",))
    col_fecha = _find_col(cols, ("fecha", "actz")) or _find_col(cols, ("fecha", "actualiz"))

    if col_id is None or col_prop is None:
        raise ValueError("No se encontraron las columnas 'ID vehículo' / 'N° de prop' en dni.xlsx.")

    df["_dni"] = df[col_id].apply(_only_digits)
    df["_prop"] = df[col_prop].apply(_only_digits)
    df = df[df["_dni"] != ""]
    df["_fecha"] = df[col_fecha] if col_fecha else None

    # Nos quedamos con el registro más reciente (fecha_actz) por DNI, priorizando
    # los que tienen número de proveedor cargado si hay empate de fecha.
    df["_fecha_sort"] = pd.to_datetime(df["_fecha"], errors="coerce")
    df["_tiene_prop"] = df["_prop"] != ""
    df = df.sort_values(["_fecha_sort", "_tiene_prop"], ascending=[False, False])
    df = df.drop_duplicates(subset="_dni", keep="first")

    rows = []
    for _, r in df.iterrows():
        fecha = r["_fecha_sort"]
        rows.append({
            "dni": r["_dni"],
            "nro_proveedor": int(r["_prop"]) if r["_prop"] else None,
            "fecha_actz": fecha.date().isoformat() if pd.notna(fecha) else None,
        })
    return db.replace_dni_proveedor(rows)


# ---------------------------------------------------------------------------
# Carga del estado real de bloqueos (foto del sistema, no de GCG)
# ---------------------------------------------------------------------------

def load_estado_bloqueos_file(file_like):
    """Excel con la lista de quiénes están bloqueados hoy de verdad (más allá de
    lo que diga la corrida de GCG). Se usa como filtro antes de consultar la API
    de GCG: si un chofer no figura acá, ya está liberado y no hace falta
    consultarlo ni generarle una tarea automática."""
    df = pd.read_excel(file_like, sheet_name=0)
    cols = df.columns.tolist()

    col_id = _find_col(cols, ("id", "vehiculo"))
    col_bloqueado = _find_col(cols, ("bloqueado",))
    col_tipo = _find_col(cols, ("tp", "veh"))
    col_desc = _find_col(cols, ("descripcion", "bloqueo"))

    if col_id is None:
        raise ValueError("No se encontró la columna 'ID vehículo' en el estado de bloqueos.")

    df["_dni"] = df[col_id].apply(_only_digits)
    df = df[df["_dni"] != ""]
    if col_bloqueado:
        df = df[df[col_bloqueado].astype(str).str.strip().str.upper() == "Y"]
    df = df.drop_duplicates(subset="_dni", keep="first")

    rows = []
    for _, r in df.iterrows():
        rows.append({
            "dni": r["_dni"],
            "tipo": r.get(col_tipo) if col_tipo else None,
            "descripcion": r.get(col_desc) if col_desc else None,
        })
    return db.replace_estado_bloqueos(rows)


# ---------------------------------------------------------------------------
# Procesamiento del exportado de GCG
# ---------------------------------------------------------------------------

def buscar_proveedor_por_dni(dni):
    """Busca el proveedor/JRT asociado a un DNI en el maestro cargado. Devuelve
    None si el DNI no está cargado en dni.xlsx o no tiene proveedor asociado."""
    dni = _only_digits(dni)
    with db.get_conn() as conn:
        row = conn.execute(
            """SELECT p.nro_proveedor, p.razon_social, p.jrt, p.mail
               FROM dni_proveedor d JOIN proveedores p ON p.nro_proveedor = d.nro_proveedor
               WHERE d.dni = ?""",
            (dni,),
        ).fetchone()
    if not row:
        return None
    return {
        "nro_proveedor": row["nro_proveedor"],
        "razon_social": row["razon_social"],
        "jrt": row["jrt"],
        "mail_proveedor": row["mail"],
    }


# ---------------------------------------------------------------------------

# Nombre de país tal como viene en la columna "País" del exportado de GCG -> código
# ISO (el mismo que devuelve la API en el campo "iso"). Claves normalizadas con _norm.
PAIS_A_ISO = {
    "argentina": "AR",
    "bolivia": "BO",
    "brasil": "BR",
    "brazil": "BR",
    "chile": "CL",
    "paraguay": "PY",
    "uruguay": "UY",
    "peru": "PE",
}


def _texto(v):
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    return str(v).strip() or None


def pais_a_iso(valor):
    """'Bolivia' -> 'BO'. Si ya viene como código de 2 letras se respeta; si es un
    país que no está en el mapeo devuelve None (se le controla todo, ver
    db.documento_aplica_a_pais)."""
    if valor is None:
        return None
    try:
        if pd.isna(valor):
            return None
    except (TypeError, ValueError):
        pass
    texto = str(valor).strip()
    if len(texto) == 2 and texto.isalpha():
        return texto.upper()
    return PAIS_A_ISO.get(_norm(texto))


EXCLUIR_CONDICION_KEYWORDS = ["extranjero"]
EXCLUIR_MATRIZ_KEYWORDS = ["independiente"]


def procesar_gcg_export(file_like, archivo_origen):
    df = pd.read_excel(file_like, sheet_name=0)
    cols = df.columns.tolist()

    col_tipo = _find_col(cols, ("tipo",))
    col_dni = _find_col(cols, ("dni",))
    col_nombre = _find_col(cols, ("nombre", "dominio")) or _find_col(cols, ("contratista",))
    col_condicion = _find_col(cols, ("condicion", "trabajador"))
    col_matriz = _find_col(cols, ("matriz", "documentos"))
    col_pais = _find_col(cols, ("pais",))
    col_contratista = next((c for c in cols if _norm(c) == "contratista"), None)
    col_situacion = next((c for c in cols if _norm(c) == "situacion"), None)

    if col_dni is None:
        raise ValueError("No se encontró la columna 'DNI' en el exportado de GCG.")

    # sólo Trabajador si la columna existe (algunos exportados no la traen)
    if col_tipo:
        df = df[df[col_tipo].astype(str).str.strip().str.lower() == "trabajador"]

    # Historial de transportes: se registra a todos los trabajadores del exportado
    # (también inactivos, dados de baja, extranjeros e independientes, que después
    # se excluyen del control).
    if col_contratista:
        fecha_obs = db.now_iso()
        db.insert_historial_contratista([
            {
                "dni": dni,
                "nombre": _texto(r.get(col_nombre)) if col_nombre else None,
                "pais": pais_a_iso(r.get(col_pais)) if col_pais else None,
                "situacion": _texto(r.get(col_situacion)) if col_situacion else None,
                "contratista": _texto(r.get(col_contratista)),
                "fuente": "exportado",
                "fecha": fecha_obs,
            }
            for _, r in df.iterrows()
            for dni in [_only_digits(r.get(col_dni))]
            if dni
        ])

    # El exportado puede traer también a los dados de baja (ya no trabajan más:
    # sirven sólo para el historial de transportes). Al control de bloqueos entran
    # los activos y los inactivos (no viajan hace unos meses, pero siguen estando).
    total_bajas = 0
    if col_situacion:
        de_baja = df[col_situacion].astype(str).map(_norm).str.contains("baja", na=False)
        total_bajas = int(de_baja.sum())
        df = df[~de_baja]

    total_evaluados_inicial = len(df)

    excluido_extranjero = pd.Series(False, index=df.index)
    if col_condicion:
        norm_cond = df[col_condicion].astype(str).map(_norm)
        excluido_extranjero = norm_cond.str.contains("extranjero", na=False)
    excluido_independiente = pd.Series(False, index=df.index)
    if col_matriz:
        norm_matriz = df[col_matriz].astype(str).map(_norm)
        excluido_independiente = norm_matriz.str.contains("independiente", na=False)

    excluidos_mask = excluido_extranjero | excluido_independiente
    total_excluidos = int(excluidos_mask.sum())
    df = df[~excluidos_mask].copy()

    documentos = db.get_documentos_control(solo_activos=True)
    doc_cols = {}
    for d in documentos:
        found = _find_col(cols, tuple(w for w in d["nombre"].split() if len(w) > 2))
        doc_cols[d["nombre"]] = found

    df["_dni"] = df[col_dni].apply(_only_digits)

    with db.get_conn() as conn:
        dni_map = {r["dni"]: r["nro_proveedor"] for r in conn.execute("SELECT dni, nro_proveedor FROM dni_proveedor")}
        prov_map = {
            r["nro_proveedor"]: dict(r)
            for r in conn.execute("SELECT * FROM proveedores")
        }

    total_sin_proveedor = 0
    out_rows = []
    for _, r in df.iterrows():
        dni = r["_dni"]
        if not dni:
            continue
        nro_prov = dni_map.get(dni)
        prov = prov_map.get(nro_prov) if nro_prov else None
        if not prov:
            total_sin_proveedor += 1
        nombre = r.get(col_nombre) if col_nombre else None
        condicion = r.get(col_condicion) if col_condicion else None
        pais = pais_a_iso(r.get(col_pais)) if col_pais else None

        for doc_nombre, doc_col in doc_cols.items():
            fecha_val = r.get(doc_col) if doc_col else None
            fecha = _parse_date(fecha_val)
            out_rows.append({
                "dni": dni,
                "nombre": nombre,
                "condicion": condicion,
                "nro_proveedor": nro_prov,
                "razon_social": prov["razon_social"] if prov else None,
                "jrt": prov["jrt"] if prov else None,
                "mail_proveedor": prov["mail"] if prov else None,
                "pais": pais,
                "documento": doc_nombre,
                "fecha_vencimiento": fecha.isoformat() if fecha else None,
            })

    run_id = db.create_run(
        fecha_run=db.now_iso(),
        archivo_origen=archivo_origen,
        total_evaluados=total_evaluados_inicial,
        total_sin_proveedor=total_sin_proveedor,
        total_excluidos=total_excluidos,
    )
    db.insert_choferes_docs(run_id, out_rows)

    missing_doc_cols = [nombre for nombre, col in doc_cols.items() if col is None]

    return {
        "run_id": run_id,
        "total_evaluados": total_evaluados_inicial,
        "total_excluidos": total_excluidos,
        "total_sin_proveedor": total_sin_proveedor,
        "total_bajas": total_bajas,
        "missing_doc_cols": missing_doc_cols,
    }


# ---------------------------------------------------------------------------
# Clasificación vencido / próximo para una corrida ya guardada
# ---------------------------------------------------------------------------

def clasificar_run(run_id, dias_alerta=7, jrt_filtro=None, orden=None, direccion="asc"):
    hoy = datetime.now().date()
    rows = db.get_choferes_docs(run_id)
    estado_bloqueos_map = db.get_estado_bloqueos_map()
    estado_bloqueos_en = db.get_config().get("estado_bloqueos_actualizado_en") or None
    # Cada documento se controla sólo para los países configurados (ej. sólo AR:
    # a los choferes extranjeros no se los bloquea por ese documento). Se filtra
    # acá y no al procesar, para que un cambio de configuración aplique sin
    # reprocesar la corrida.
    docs_config = {d["nombre"]: d for d in db.get_documentos_control()}

    por_chofer = {}
    for r in rows:
        jrt_label = (r["jrt"] or "").strip() or SIN_JRT_LABEL
        if jrt_filtro is not None and jrt_label not in jrt_filtro:
            continue
        doc_config = docs_config.get(r["documento"])
        if doc_config and not db.documento_aplica_a_pais(doc_config, r.get("pais")):
            continue
        key = r["dni"]
        jd_info = estado_bloqueos_map.get(key)
        c = por_chofer.setdefault(key, {
            "dni": r["dni"],
            "nombre": r["nombre"],
            "condicion": r["condicion"],
            "pais": r.get("pais"),
            "nro_proveedor": r["nro_proveedor"],
            "razon_social": r["razon_social"],
            "jrt": jrt_label,
            "mail_proveedor": r["mail_proveedor"],
            "docs_vencidos": [],
            "docs_proximos": [],
            "docs_todos": [],
            "sin_proveedor": r["nro_proveedor"] is None,
            "jd_bloqueado": jd_info is not None,
            "jd_tipo": (jd_info or {}).get("tipo"),
            "jd_descripcion": (jd_info or {}).get("descripcion"),
        })
        fecha_str = r["fecha_vencimiento"]
        if not fecha_str:
            # Sin fecha cargada en GCG para este documento: no se evalúa (igual que
            # el criterio manual de referencia), no bloquea por sí solo.
            c["docs_todos"].append({"documento": r["documento"], "fecha": None, "dias": None, "estado": "sin_dato"})
            continue
        fecha = datetime.fromisoformat(fecha_str).date()
        dias_restantes = (fecha - hoy).days
        if dias_restantes < 0:
            estado_doc = "vencido"
            c["docs_vencidos"].append({"documento": r["documento"], "fecha": fecha.isoformat(), "dias": dias_restantes})
        elif dias_restantes <= dias_alerta:
            estado_doc = "proximo"
            c["docs_proximos"].append({"documento": r["documento"], "fecha": fecha.isoformat(), "dias": dias_restantes})
        else:
            estado_doc = "vigente"
        c["docs_todos"].append({
            "documento": r["documento"], "fecha": fecha.isoformat(), "dias": dias_restantes, "estado": estado_doc,
        })

    vencidos = []
    proximos = []
    for c in por_chofer.values():
        c["docs_todos"].sort(key=lambda d: d["documento"])
        if c["docs_vencidos"]:
            vencidos.append(c)
        elif c["docs_proximos"]:
            proximos.append(c)

    vencidos.sort(key=lambda c: ((c["jrt"] or ""), c["nombre"] or ""))
    proximos.sort(key=lambda c: (min((d["dias"] for d in c["docs_proximos"]), default=999), (c["jrt"] or "")))

    if orden:
        vencidos = _ordenar_choferes(vencidos, orden, direccion)
        proximos = _ordenar_choferes(proximos, orden, direccion)

    return {
        "vencidos": vencidos,
        "proximos": proximos,
        "hoy": hoy.isoformat(),
        "dias_alerta": dias_alerta,
        "resumen_jrt_vencidos": _resumen_por_jrt(vencidos),
        "resumen_jrt_proximos": _resumen_por_jrt(proximos),
        "estado_bloqueos_en": estado_bloqueos_en,
        "estado_bloqueos_cargado": bool(estado_bloqueos_en),
    }


_ORDEN_CAMPOS = {
    "jrt": lambda c: (c["jrt"] or ""),
    "dni": lambda c: (c["dni"] or ""),
    "nombre": lambda c: (c["nombre"] or ""),
    "proveedor": lambda c: (c["razon_social"] or ""),
}


def _ordenar_choferes(lista, campo, direccion):
    key_fn = _ORDEN_CAMPOS.get(campo)
    if not key_fn:
        return lista
    return sorted(lista, key=key_fn, reverse=(direccion == "desc"))


def _resumen_por_jrt(lista):
    conteo = {}
    for c in lista:
        conteo[c["jrt"]] = conteo.get(c["jrt"], 0) + 1
    # Por defecto, el JRT con más casos va arriba (empate: alfabético). El
    # cuadrito también se puede reordenar a mano desde la tabla (JS, sin reload).
    filas = sorted(conteo.items(), key=lambda kv: (-kv[1], kv[0]))
    total = sum(conteo.values())
    return {"filas": filas, "total": total}


def jrt_list_for_run(run_id):
    rows = db.get_choferes_docs(run_id)
    jrts = sorted({(r["jrt"] or SIN_JRT_LABEL) for r in rows})
    return jrts


# ---------------------------------------------------------------------------
# Agrupado por transporte (proveedor), para el envío de mails
# ---------------------------------------------------------------------------

def resumen_por_transporte(run_id, dias_alerta=7, jrt_filtro=None):
    resultado = clasificar_run(run_id, dias_alerta=dias_alerta, jrt_filtro=jrt_filtro)
    transportes = {}

    def _agregar(lista, campo):
        for c in lista:
            key = c["nro_proveedor"] if c["nro_proveedor"] else f"sin__{c['razon_social'] or c['jrt']}"
            t = transportes.setdefault(key, {
                "nro_proveedor": c["nro_proveedor"],
                "razon_social": c["razon_social"] or "Sin proveedor",
                "jrt": c["jrt"],
                "mail_proveedor": c["mail_proveedor"],
                "vencidos": [],
                "proximos": [],
            })
            t[campo].append(c)

    _agregar(resultado["vencidos"], "vencidos")
    _agregar(resultado["proximos"], "proximos")

    lista = list(transportes.values())
    lista.sort(key=lambda t: ((t["jrt"] or ""), t["razon_social"] or ""))
    return lista


def exportar_excel_transporte(transporte):
    """Genera un .xlsx con dos hojas (Vencidos / Próximos) para un transporte,
    con una columna por documento controlado (igual que en el dashboard)."""
    documentos_activos = [d["nombre"] for d in db.get_documentos_control(solo_activos=True)]
    columnas_base = ["Chofer", "DNI", "JRT", "N° Proveedor", "Proveedor"]
    columnas = columnas_base + documentos_activos
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for campo, hoja in (("vencidos", "Vencidos"), ("proximos", "Proximos a vencer")):
            filas, estados = _filas_por_documento(transporte[campo], documentos_activos)
            df = pd.DataFrame(filas, columns=columnas)
            df.to_excel(writer, index=False, sheet_name=hoja)
            _pintar_columnas_documentos(writer.sheets[hoja], len(columnas_base), estados)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# Exportar a Excel
# ---------------------------------------------------------------------------

_FILL_POR_ESTADO = {
    "vencido": "FDECEA",
    "proximo": "FEF3D9",
    "vigente": "E6F4EA",
    "sin_dato": "EEF0F2",
}


def _motivo_jd(c):
    partes = [
        str(p).strip() for p in (c.get("jd_tipo"), c.get("jd_descripcion"))
        if p not in (None, "") and str(p).strip()
    ]
    return " — ".join(partes)


def _filas_por_documento(choferes, documentos_activos, incluir_jd=False, jd_cargado=True):
    """Arma filas con una columna por documento (fecha o 'sin dato') a partir de
    docs_todos, y en paralelo la lista de estados de esas columnas para pintarlas."""
    filas = []
    estados = []
    for c in choferes:
        docs_por_nombre = {d["documento"]: d for d in c["docs_todos"]}
        fila = {
            "Chofer": c["nombre"],
            "DNI": c["dni"],
            "JRT": c["jrt"],
            "N° Proveedor": c["nro_proveedor"],
            "Proveedor": c["razon_social"],
        }
        if incluir_jd:
            if not jd_cargado:
                fila["Situación JD"] = "Sin datos"
                fila["Motivo JD"] = ""
            else:
                fila["Situación JD"] = "Bloqueado" if c.get("jd_bloqueado") else "No bloqueado"
                fila["Motivo JD"] = _motivo_jd(c) if c.get("jd_bloqueado") else ""
        fila_estados = []
        for doc_nombre in documentos_activos:
            d = docs_por_nombre.get(doc_nombre)
            fila[doc_nombre] = (db.fmt_fecha(d["fecha"]) if d["fecha"] else "sin dato") if d else "—"
            fila_estados.append(d["estado"] if d else None)
        filas.append(fila)
        estados.append(fila_estados)
    return filas, estados


def _pintar_columnas_documentos(worksheet, offset, estados_por_fila):
    """offset: cantidad de columnas base antes de las columnas de documentos."""
    from openpyxl.styles import PatternFill

    for row_idx, fila_estados in enumerate(estados_por_fila, start=2):  # fila 1 = encabezado
        for col_idx, estado in enumerate(fila_estados, start=offset + 1):
            color = _FILL_POR_ESTADO.get(estado)
            if color:
                worksheet.cell(row=row_idx, column=col_idx).fill = PatternFill(
                    start_color=color, end_color=color, fill_type="solid"
                )


def exportar_excel(lista):
    """Cada documento controlado tiene su propia columna con la fecha (o 'sin dato'),
    igual que en el dashboard. Devuelve un .xlsx."""
    documentos_activos = [d["nombre"] for d in db.get_documentos_control(solo_activos=True)]
    jd_cargado = bool(db.get_config().get("estado_bloqueos_actualizado_en"))
    columnas_base = ["Chofer", "DNI", "JRT", "N° Proveedor", "Proveedor", "Situación JD", "Motivo JD"]
    columnas = columnas_base + documentos_activos

    filas, estados = _filas_por_documento(lista, documentos_activos, incluir_jd=True, jd_cargado=jd_cargado)
    df = pd.DataFrame(filas, columns=columnas)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Listado")
        _pintar_columnas_documentos(writer.sheets["Listado"], len(columnas_base), estados)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# Historial de transportes (en qué contratista trabajó cada chofer)
# ---------------------------------------------------------------------------

FUENTE_HIST_LABELS = {
    "exportado": "Exportado GCG",
    "mapeo_mensual": "Barrido mensual API",
    "solicitud": "Solicitud de desbloqueo",
    "consulta": "Buscador GCG",
}


def periodos_contratista(observaciones):
    """Períodos de UN chofer en cada contratista, a partir de sus observaciones
    ordenadas por fecha: [{contratista, cuit, desde, hasta, veces, fuentes,
    actual}], ordenados por "desde".

    Cada "foto" (mismo momento y fuente) puede traer más de un contratista: en el
    exportado un chofer puede figurar en dos transportes a la vez, y eso no es un
    cambio. Un período sigue abierto mientras el contratista aparezca; lo cierra
    una foto del exportado en la que ya no figura, o una de la API (que trae un
    solo contratista) si el chofer estaba en un único transporte."""
    fotos = []
    for o in observaciones:
        if fotos and fotos[-1]["fecha"] == o["fecha"] and fotos[-1]["fuente"] == o["fuente"]:
            fotos[-1]["obs"].append(o)
        else:
            fotos.append({"fecha": o["fecha"], "fuente": o["fuente"], "obs": [o]})

    periodos = []
    abiertos = {}  # clave -> período
    for foto in fotos:
        claves_foto = {}
        for o in foto["obs"]:
            claves_foto.setdefault(_norm(o["contratista"]) or "(sin contratista)", o)
        cierra = foto["fuente"] == "exportado" or len(abiertos) <= 1
        if cierra:
            for clave in [c for c in abiertos if c not in claves_foto]:
                del abiertos[clave]
        fuente_label = FUENTE_HIST_LABELS.get(foto["fuente"], foto["fuente"])
        for clave, o in claves_foto.items():
            p = abiertos.get(clave)
            if p is None:
                p = {
                    "_clave": clave,
                    "contratista": o["contratista"] or "(sin contratista)",
                    "cuit": None,
                    "desde": foto["fecha"],
                    "hasta": foto["fecha"],
                    "veces": 0,
                    "fuentes": set(),
                    "situacion": None,
                }
                periodos.append(p)
                abiertos[clave] = p
            p["hasta"] = foto["fecha"]
            p["veces"] += 1
            p["cuit"] = p["cuit"] or o.get("contratista_cuit")
            p["situacion"] = o.get("situacion") or p["situacion"]
            p["fuentes"].add(fuente_label)

    for p in periodos:
        p["actual"] = abiertos.get(p["_clave"]) is p
        p["fuentes"] = ", ".join(sorted(p["fuentes"]))
    periodos.sort(key=lambda p: (p["desde"], p["contratista"]))
    return periodos


def _ancho_columnas(ws, df):
    for col_idx, nombre_col in enumerate(df.columns, start=1):
        largo = df[nombre_col].map(lambda v: len(str(v)) if v is not None and v == v else 0).max() if len(df) else 0
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max(len(str(nombre_col)), 10, largo) + 2, 50)


def exportar_excel_transportes():
    """Excel con tres hojas: dónde figura hoy cada chofer, quiénes cambiaron de
    transporte, y el historial completo por períodos."""
    from openpyxl.styles import Font

    por_dni = {}
    for o in db.get_historial_contratista():
        por_dni.setdefault(o["dni"], []).append(o)

    actual, cambios, historial = [], [], []
    for dni, obs in por_dni.items():
        nombre = next((o["nombre"] for o in reversed(obs) if o["nombre"]), None)
        pais = next((o["pais"] for o in reversed(obs) if o["pais"]), None)
        periodos = periodos_contratista(obs)
        vigentes = [p for p in periodos if p["actual"]]
        cerrados = [p for p in periodos if not p["actual"]]
        actual.append({
            "Chofer": nombre, "DNI": dni, "País": pais,
            "Contratista actual": " / ".join(p["contratista"] for p in vigentes),
            "Situación": " / ".join(p["situacion"] or "—" for p in vigentes),
            "CUIT contratista": " / ".join(p["cuit"] for p in vigentes if p["cuit"]),
            "En este transporte desde": db.fmt_fecha(min(p["desde"] for p in vigentes)) if vigentes else "",
            "Visto por última vez": db.fmt_fecha(obs[-1]["fecha"]),
            "Transportes distintos": len({p["_clave"] for p in periodos}),
        })
        if cerrados:
            anterior = max(cerrados, key=lambda p: p["hasta"])
            cambios.append({
                "Chofer": nombre, "DNI": dni, "País": pais,
                "Transporte anterior": anterior["contratista"],
                "Visto ahí hasta": db.fmt_fecha(anterior["hasta"]),
                "Transporte actual": " / ".join(p["contratista"] for p in vigentes),
                "Visto ahí desde": db.fmt_fecha(min(p["desde"] for p in vigentes)) if vigentes else "",
                "Cambios": len(cerrados),
            })
        for p in periodos:
            historial.append({
                "Chofer": nombre, "DNI": dni, "Contratista": p["contratista"], "CUIT contratista": p["cuit"],
                "Desde": db.fmt_fecha(p["desde"]), "Hasta": db.fmt_fecha(p["hasta"]),
                "Actual": "Sí" if p["actual"] else "No",
                "Situación": p["situacion"],
                "Veces visto": p["veces"], "Fuentes": p["fuentes"],
            })

    hojas = (
        ("Donde trabaja hoy", sorted(actual, key=lambda f: ((f["Contratista actual"] or ""), f["Chofer"] or ""))),
        ("Cambios de transporte", sorted(cambios, key=lambda f: (f["Chofer"] or ""))),
        ("Historial completo", historial),
    )
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for nombre_hoja, filas in hojas:
            df = pd.DataFrame(filas)
            if df.empty:
                df = pd.DataFrame({"Sin datos": []})
            df.to_excel(writer, index=False, sheet_name=nombre_hoja)
            ws = writer.sheets[nombre_hoja]
            for cell in ws[1]:
                cell.font = Font(bold=True)
            _ancho_columnas(ws, df)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# Exportar solicitudes de desbloqueo (auditoría)
# ---------------------------------------------------------------------------

ORIGEN_LABELS = {
    "manual": "Manual",
    "jrt_auto": "Auto (JRT)",
    "admin_auto": "Auto (Admin)",
    "gcg_auto": "Auto (GCG)",
    "criticos_verdes": "Auto (críticos verdes)",
}

_FILL_POR_ESTADO_SOLICITUD = {
    "solicitado": "FEF3D9",
    "autorizado": "DBEAFE",
    "rechazado": "FDECEA",
    "ejecutado": "E6F4EA",
}


def _documentos_por_dni(run_id):
    """dni -> {documento: fecha_vencimiento (iso) o None} para una corrida. Se usa
    sólo como resguardo para solicitudes viejas, generadas antes de que cada
    solicitud guardara su propia consulta a GCG."""
    if not run_id:
        return {}
    out = {}
    for r in db.get_choferes_docs(run_id):
        out.setdefault(r["dni"], {})[r["documento"]] = r["fecha_vencimiento"]
    return out


def _documentos_desde_gcg_json(gcg_json, documentos_control):
    """A partir del JSON de GCG guardado en la propia solicitud (la foto real de
    lo que se controló en ese momento), arma {documento: {fecha, estado, encontrado}}
    usando el mismo criterio (Nº GCG) que la consulta en vivo. None si no hay JSON
    guardado (solicitud vieja, o falló la consulta)."""
    if not gcg_json:
        return None
    try:
        data = json.loads(gcg_json)
    except (ValueError, TypeError):
        return None
    return {d["nombre"]: d for d in gcg_api.evaluar_criticos(data, documentos_control)}


def exportar_excel_solicitudes(solicitudes):
    """Auditoría completa de solicitudes de desbloqueo: quién la pidió, quién la
    autorizó (o rechazó) y quién la ejecutó, con fecha y hora de cada paso, y los
    documentos tal como se vieron en GCG al momento de esa solicitud (una columna
    por documento con su fecha), en una columna por documento con su fecha (igual
    que en el resto de los exportados)."""
    from openpyxl.styles import Font, PatternFill

    documentos_control = db.get_documentos_control(solo_activos=True)
    documentos_activos = [d["nombre"] for d in documentos_control]
    columnas_base = ["ID", "Chofer", "DNI", "JRT", "Proveedor", "N° Proveedor"]
    columnas_auditoria = [
        "Origen", "Estado", "Comentario",
        "Solicitado por", "Fecha solicitud",
        "Autorizado por", "Fecha autorización",
        "Ejecutado por", "Fecha ejecución",
        "Consulta a GCG",
        "Corrida (run_id)",
    ]
    columnas = columnas_base + documentos_activos + columnas_auditoria

    hoy = datetime.now().date()
    cache_docs_por_run = {}
    filas = []
    estados_solicitud = []
    estados_doc_filas = []
    for s in solicitudes:
        docs_gcg = _documentos_desde_gcg_json(s.get("gcg_json"), documentos_control)

        fila = {
            "ID": s["id"],
            "Chofer": s["nombre"],
            "DNI": s["dni"],
            "JRT": s["jrt"],
            "Proveedor": s["razon_social"],
            "N° Proveedor": s["nro_proveedor"],
        }
        fila_estados_doc = []
        if docs_gcg is not None:
            # Foto real controlada por API al momento de esta solicitud puntual.
            for doc_nombre in documentos_activos:
                d = docs_gcg.get(doc_nombre)
                if d is None:
                    # No se le controla a un chofer de ese país.
                    fila[doc_nombre] = "no aplica"
                    fila_estados_doc.append(None)
                elif d["encontrado"]:
                    fila[doc_nombre] = d["fecha"] or "sin dato"
                    fila_estados_doc.append("vigente" if d["estado"] else "vencido")
                else:
                    fila[doc_nombre] = "no encontrado en GCG"
                    fila_estados_doc.append(None)
        else:
            # Resguardo: solicitudes generadas antes de guardar la consulta propia,
            # o a las que les falló la consulta a GCG. Se usan los datos de la
            # corrida cargada (pueden no coincidir con la fecha real del control).
            run_id = s.get("run_id")
            if run_id not in cache_docs_por_run:
                cache_docs_por_run[run_id] = _documentos_por_dni(run_id)
            docs_dni = cache_docs_por_run[run_id].get(s["dni"], {})
            for doc_nombre in documentos_activos:
                if doc_nombre in docs_dni:
                    fecha_str = docs_dni[doc_nombre]
                    if fecha_str:
                        fecha = datetime.fromisoformat(fecha_str).date()
                        fila[doc_nombre] = db.fmt_fecha(fecha_str)
                        fila_estados_doc.append("vencido" if fecha < hoy else "vigente")
                    else:
                        fila[doc_nombre] = "sin dato"
                        fila_estados_doc.append("sin_dato")
                else:
                    fila[doc_nombre] = "—"
                    fila_estados_doc.append(None)

        fila.update({
            "Origen": ORIGEN_LABELS.get(s.get("origen"), s.get("origen") or "Manual"),
            "Estado": (s["estado"] or "").capitalize(),
            "Comentario": s.get("comentario") or "",
            "Solicitado por": s.get("solicitado_por_nombre") or "",
            "Fecha solicitud": db.fmt_fecha(s["fecha_solicitud"], con_hora=True) if s.get("fecha_solicitud") else "",
            "Autorizado por": s.get("autorizado_por_nombre") or "",
            "Fecha autorización": db.fmt_fecha(s["fecha_autorizacion"], con_hora=True) if s.get("fecha_autorizacion") else "",
            "Ejecutado por": s.get("ejecutado_por_nombre") or "",
            "Fecha ejecución": db.fmt_fecha(s["fecha_ejecucion"], con_hora=True) if s.get("fecha_ejecucion") else "",
            "Consulta a GCG": (
                db.fmt_fecha(s["gcg_consultado_en"], con_hora=True) if s.get("gcg_consultado_en") and not s.get("gcg_error")
                else (f"Error: {s['gcg_error']}" if s.get("gcg_error") else "—")
            ),
            "Corrida (run_id)": s.get("run_id"),
        })
        filas.append(fila)
        estados_solicitud.append(s.get("estado"))
        estados_doc_filas.append(fila_estados_doc)

    df = pd.DataFrame(filas, columns=columnas)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Solicitudes")
        ws = writer.sheets["Solicitudes"]

        for cell in ws[1]:
            cell.font = Font(bold=True)

        col_estado = columnas.index("Estado") + 1
        for row_idx, estado in enumerate(estados_solicitud, start=2):
            color = _FILL_POR_ESTADO_SOLICITUD.get(estado)
            if color:
                ws.cell(row=row_idx, column=col_estado).fill = PatternFill(
                    start_color=color, end_color=color, fill_type="solid"
                )

        _pintar_columnas_documentos(ws, len(columnas_base), estados_doc_filas)

        for col_idx, nombre_col in enumerate(columnas, start=1):
            ancho = min(max(len(nombre_col), 12, df[nombre_col].astype(str).map(len).max() if len(df) else 0) + 2, 60)
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = ancho

    buf.seek(0)
    return buf
